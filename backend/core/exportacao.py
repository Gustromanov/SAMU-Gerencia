"""Geradores de relatório em PDF e XLSX para o SAMU 192."""
from __future__ import annotations

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import models


def _coletar_dados() -> dict:
    cfg = models.Configuracao.obter()
    lotes = list(models.Lote.objects.select_related(
        "medicamento", "medicamento__tipo", "medicamento__forma"))
    veiculos = list(models.Veiculo.objects.all())
    funcionarios = list(models.Funcionario.objects.all())
    hoje = timezone.now().date()

    kpis = {
        "medicamentos": models.Medicamento.objects.count(),
        "lotes": len(lotes),
        "estoque_total": sum(l.quantidade for l in lotes),
        "baixo": sum(1 for l in lotes if l.nivel_estoque == "BAIXO"),
        "muito_baixo": sum(1 for l in lotes if l.nivel_estoque == "MUITO_BAIXO"),
        "vencendo": sum(1 for l in lotes if l.nivel_validade in {"VENCIDO", "MUITO_PROXIMO"}),
        "funcionarios": len(funcionarios),
        "veiculos": len(veiculos),
    }
    por_tipo, por_forma = {}, {}
    for l in lotes:
        por_tipo[l.medicamento.tipo.nome] = por_tipo.get(l.medicamento.tipo.nome, 0) + l.quantidade
        por_forma[l.medicamento.forma.nome] = por_forma.get(l.medicamento.forma.nome, 0) + l.quantidade
    frota_status = {"NORMAL": 0, "MANUTENCAO": 0, "INATIVO": 0}
    frota_tipo = {"VAN": 0, "CARRO": 0, "MOTO": 0}
    for v in veiculos:
        frota_status[v.status] = frota_status.get(v.status, 0) + 1
        frota_tipo[v.tipo] = frota_tipo.get(v.tipo, 0) + 1
    situacao = {"RECEBIDO": 0, "LOTEADO": 0, "DOADO": 0}
    for l in lotes:
        situacao[l.situacao] = situacao.get(l.situacao, 0) + l.quantidade
    movimentacoes = list(models.Movimentacao.objects.select_related(
        "lote", "lote__medicamento", "usuario", "funcionario").order_by("-horario")[:200])
    lotes.sort(key=lambda l: l.validade)
    return {
        "kpis": kpis, "cfg": cfg,
        "por_tipo": sorted(por_tipo.items(), key=lambda x: -x[1]),
        "por_forma": sorted(por_forma.items(), key=lambda x: -x[1]),
        "frota_status": frota_status, "frota_tipo": frota_tipo, "situacao": situacao,
        "muito_baixo": [l for l in lotes if l.nivel_estoque == "MUITO_BAIXO"],
        "baixo": [l for l in lotes if l.nivel_estoque == "BAIXO"],
        "vencendo": [l for l in lotes if l.nivel_validade in {"VENCIDO", "MUITO_PROXIMO"}],
        "controlados": [l for l in lotes if l.nivel_validade == "CONTROLADO"],
        "lotes": lotes, "funcionarios": funcionarios, "veiculos": veiculos,
        "movimentacoes": movimentacoes, "hoje": hoje,
    }


def gerar_xlsx() -> bytes:
    d = _coletar_dados()
    wb = Workbook()

    hfont = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor="7F1D1D")
    tfont = Font(bold=True, size=16, color="7F1D1D")
    center = Alignment(horizontal="center", vertical="center")

    def header(ws, row, labels):
        for i, h in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=h); c.font = hfont; c.fill = hfill; c.alignment = center

    # Resumo
    ws = wb.active; ws.title = "Resumo"
    ws["A1"] = "SAMU 192 · Relatório operacional"; ws["A1"].font = tfont
    ws.merge_cells("A1:D1")
    ws["A2"] = f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
    ws.merge_cells("A2:D2")
    header(ws, 4, ["Indicador", "Valor"])
    k = d["kpis"]
    for i, (lbl, v) in enumerate([
        ("Medicamentos", k["medicamentos"]), ("Lotes", k["lotes"]),
        ("Estoque total (un)", k["estoque_total"]),
        ("Lotes em nível baixo", k["baixo"]), ("Lotes muito baixos", k["muito_baixo"]),
        ("Vencendo / Vencidos", k["vencendo"]),
        ("Funcionários", k["funcionarios"]), ("Veículos", k["veiculos"]),
    ], start=5):
        ws.cell(row=i, column=1, value=lbl); ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 36; ws.column_dimensions["B"].width = 16

    # Medicamentos (todos)
    ws2 = wb.create_sheet("Medicamentos")
    header(ws2, 1, ["Medicamento", "Metragem", "Tipo", "Forma", "Complemento", "Lotes", "Qtd. total", "Criado em"])
    for i, m in enumerate(models.Medicamento.objects.select_related("tipo", "forma").all(), start=2):
        ws2.cell(row=i, column=1, value=m.nome); ws2.cell(row=i, column=2, value=m.metragem)
        ws2.cell(row=i, column=3, value=m.tipo.nome); ws2.cell(row=i, column=4, value=m.forma.nome)
        ws2.cell(row=i, column=5, value=m.complemento); ws2.cell(row=i, column=6, value=m.lotes.count())
        ws2.cell(row=i, column=7, value=m.quantidade_total)
        ws2.cell(row=i, column=8, value=timezone.localtime(m.criado_em).strftime("%d/%m/%Y %H:%M"))
    for i, w in enumerate([30, 14, 20, 18, 32, 8, 12, 18], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Lotes (todos)
    ws3 = wb.create_sheet("Lotes")
    header(ws3, 1, ["Medicamento", "Lote", "Tipo", "Forma", "Qtd.", "Cadastrado", "Validade", "Dias restantes", "Situação", "Nível", "Validade (classif.)"])
    for i, l in enumerate(d["lotes"], start=2):
        ws3.cell(row=i, column=1, value=l.medicamento.nome_exibicao)
        ws3.cell(row=i, column=2, value=l.codigo); ws3.cell(row=i, column=3, value=l.medicamento.tipo.nome)
        ws3.cell(row=i, column=4, value=l.medicamento.forma.nome); ws3.cell(row=i, column=5, value=l.quantidade)
        ws3.cell(row=i, column=6, value=timezone.localtime(l.criado_em).strftime("%d/%m/%Y"))
        ws3.cell(row=i, column=7, value=l.validade.strftime("%d/%m/%Y"))
        ws3.cell(row=i, column=8, value=l.dias_para_validade)
        ws3.cell(row=i, column=9, value=l.get_situacao_display())
        ws3.cell(row=i, column=10, value=l.nivel_estoque_label)
        ws3.cell(row=i, column=11, value=l.nivel_validade_label)
    for i, w in enumerate([30, 14, 18, 16, 8, 14, 14, 14, 14, 14, 18], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # Distribuições
    ws4 = wb.create_sheet("Distribuições")
    row = 1
    for titulo, pares in [
        ("Estoque por tipo", d["por_tipo"]), ("Estoque por forma", d["por_forma"]),
        ("Situação dos lotes", [("Recebidos", d["situacao"]["RECEBIDO"]),
                                 ("Entregues", d["situacao"]["LOTEADO"]),
                                 ("Doados", d["situacao"]["DOADO"])]),
        ("Frota por status", [("Normal", d["frota_status"]["NORMAL"]),
                               ("Manutenção", d["frota_status"]["MANUTENCAO"]),
                               ("Inativo", d["frota_status"]["INATIVO"])]),
        ("Frota por tipo", [("Van", d["frota_tipo"]["VAN"]),
                             ("Carro", d["frota_tipo"]["CARRO"]),
                             ("Motolância", d["frota_tipo"]["MOTO"])]),
    ]:
        ws4.cell(row=row, column=1, value=titulo).font = Font(bold=True, color="7F1D1D", size=13)
        row += 1
        header(ws4, row, ["Categoria", "Valor"]); row += 1
        total = sum(v for _, v in pares) or 1
        for cat, val in pares:
            ws4.cell(row=row, column=1, value=cat); ws4.cell(row=row, column=2, value=val)
            ws4.cell(row=row, column=3, value=f"{(val / total * 100):.1f}%")
            row += 1
        row += 1
    ws4.column_dimensions["A"].width = 30; ws4.column_dimensions["B"].width = 14; ws4.column_dimensions["C"].width = 10

    # Funcionários
    ws5 = wb.create_sheet("Funcionários")
    header(ws5, 1, ["Nome", "CPF", "Função", "E-mail", "Telefone"])
    for i, f in enumerate(d["funcionarios"], start=2):
        ws5.cell(row=i, column=1, value=f.nome); ws5.cell(row=i, column=2, value=f.cpf)
        ws5.cell(row=i, column=3, value=f.get_funcao_display())
        ws5.cell(row=i, column=4, value=f.email); ws5.cell(row=i, column=5, value=f.telefone)
    for i, w in enumerate([30, 16, 24, 30, 18], start=1):
        ws5.column_dimensions[get_column_letter(i)].width = w

    # Veículos
    ws6 = wb.create_sheet("Veículos")
    header(ws6, 1, ["ID", "Modelo", "Ano", "Tipo", "Status"])
    for i, v in enumerate(d["veiculos"], start=2):
        ws6.cell(row=i, column=1, value=v.codigo_id); ws6.cell(row=i, column=2, value=v.modelo)
        ws6.cell(row=i, column=3, value=v.ano); ws6.cell(row=i, column=4, value=v.get_tipo_display())
        ws6.cell(row=i, column=5, value=v.get_status_display())
    for i, w in enumerate([14, 32, 8, 18, 16], start=1):
        ws6.column_dimensions[get_column_letter(i)].width = w

    # Histórico (movimentações)
    ws7 = wb.create_sheet("Histórico")
    header(ws7, 1, ["Horário", "Usuário", "Tipo", "Medicamento", "Lote", "Quantidade", "Origem", "Destino"])
    for i, m in enumerate(d["movimentacoes"], start=2):
        local = timezone.localtime(m.horario)
        ws7.cell(row=i, column=1, value=local.strftime("%d/%m/%Y %H:%M"))
        ws7.cell(row=i, column=2, value=(m.usuario.get_full_name() if m.usuario else "-") or (m.funcionario.nome if m.funcionario else "-"))
        ws7.cell(row=i, column=3, value=m.get_tipo_display())
        ws7.cell(row=i, column=4, value=m.lote.medicamento.nome_exibicao)
        ws7.cell(row=i, column=5, value=m.lote.codigo); ws7.cell(row=i, column=6, value=m.quantidade)
        ws7.cell(row=i, column=7, value=m.origem); ws7.cell(row=i, column=8, value=m.destino)
    for i, w in enumerate([18, 24, 10, 30, 14, 10, 18, 18], start=1):
        ws7.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO(); wb.save(buf); return buf.getvalue()


def gerar_pdf() -> bytes:
    d = _coletar_dados()
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title="Relatório SAMU 192", author="SAMU 192")
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle("titulo", parent=styles["Title"], fontSize=22,
                            textColor=colors.HexColor("#7F1D1D"), alignment=0, spaceAfter=4, leading=26)
    subtitulo = ParagraphStyle("sub", parent=styles["Normal"], fontSize=10,
                               textColor=colors.HexColor("#525252"), spaceAfter=14)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13,
                        textColor=colors.HexColor("#7F1D1D"), spaceBefore=10, spaceAfter=6)
    note = ParagraphStyle("note", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#525252"))

    el = []
    el.append(Paragraph("SAMU 192 · Relatório operacional", titulo))
    el.append(Paragraph(f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')} · Unidade local", subtitulo))

    k = d["kpis"]
    kpi_data = [
        ["Medicamentos", "Lotes", "Estoque", "Baixo", "Muito baixo", "Vencendo", "Funcionários", "Veículos"],
        [k["medicamentos"], k["lotes"], k["estoque_total"], k["baixo"], k["muito_baixo"], k["vencendo"], k["funcionarios"], k["veiculos"]],
    ]
    t = Table(kpi_data, colWidths=[32 * mm] * 8, rowHeights=[8 * mm, 14 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FEF2F2")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#7F1D1D")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9), ("FONTSIZE", (0, 1), (-1, 1), 18),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E5E5")),
    ]))
    el.append(t); el.append(Spacer(1, 6 * mm))

    # Medicamentos (resumo com data cadastro)
    el.append(Paragraph("Medicamentos cadastrados (data de cadastro)", h2))
    med_data = [["Medicamento", "Metragem", "Tipo", "Forma", "Lotes", "Qtd.", "Cadastrado em"]]
    for m in models.Medicamento.objects.select_related("tipo", "forma")[:30]:
        med_data.append([m.nome, m.metragem, m.tipo.nome, m.forma.nome,
                         m.lotes.count(), m.quantidade_total,
                         timezone.localtime(m.criado_em).strftime("%d/%m/%Y")])
    t = Table(med_data, colWidths=[55 * mm, 20 * mm, 32 * mm, 28 * mm, 15 * mm, 18 * mm, 28 * mm], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7F1D1D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D4D4D4")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
    ]))
    el.append(t); el.append(Spacer(1, 4 * mm))

    # Lotes detalhados (com validade, dias)
    el.append(Paragraph("Lotes com data de validade e dias restantes", h2))
    lot_data = [["Medicamento", "Lote", "Qtd", "Validade", "Dias", "Nível", "Classif."]]
    for l in d["lotes"][:30]:
        lot_data.append([l.medicamento.nome_exibicao, l.codigo, l.quantidade,
                         l.validade.strftime("%d/%m/%Y"), l.dias_para_validade,
                         l.nivel_estoque_label, l.nivel_validade_label])
    t = Table(lot_data, colWidths=[60 * mm, 22 * mm, 15 * mm, 25 * mm, 15 * mm, 25 * mm, 25 * mm], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C2410C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D4D4D4")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFEDD5")]),
    ]))
    el.append(t); el.append(Spacer(1, 4 * mm))

    # Histórico recente
    el.append(Paragraph("Histórico recente de manipulações (todos os usuários)", h2))
    hist = [["Horário", "Usuário", "Tipo", "Medicamento", "Lote", "Qtd"]]
    for m in d["movimentacoes"][:30]:
        usuario = (m.usuario.get_full_name() if m.usuario else None) or (m.funcionario.nome if m.funcionario else "-")
        hist.append([timezone.localtime(m.horario).strftime("%d/%m/%Y %H:%M"), usuario,
                     m.get_tipo_display(), m.lote.medicamento.nome_exibicao, m.lote.codigo, m.quantidade])
    t = Table(hist, colWidths=[32 * mm, 42 * mm, 18 * mm, 60 * mm, 25 * mm, 15 * mm], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#525252")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D4D4D4")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
    ]))
    el.append(t)

    doc.build(el)
    return buf.getvalue()
